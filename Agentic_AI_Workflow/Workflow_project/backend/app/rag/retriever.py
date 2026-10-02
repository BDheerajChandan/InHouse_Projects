"""
app/rag/retriever.py
Dynamic document retrieval pipeline.
- OpenAI embeddings are initialised lazily (only when retrieve() is called).
- Supports local files, multiple files, folders, and non-local sources.
- Source type is determined at runtime from node config.
- API key: Static mode -> .env (OPENAI_API_KEY); Dynamic mode -> key passed from node config.
"""

from __future__ import annotations

import os
import tempfile
from pathlib import Path
from typing import Any

import requests
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

BASE_DIR = Path(__file__).resolve().parent.parent.parent
CHROMA_DIR = BASE_DIR / "Chroma_db"

# Supported extensions for local folder/multi-file discovery
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".html", ".htm", ".md"}


# ─── Lazy embeddings ──────────────────────────────────────────────────────────

def _get_embeddings(api_key: str | None = None):
    """
    Return OpenAI embeddings instance. Raises with a clear message on failure.
    If api_key is provided (Dynamic mode) it is used; otherwise falls back to .env (Static mode).
    """
    from langchain_openai import OpenAIEmbeddings
    key = (api_key or "").strip() or os.getenv("OPENAI_API_KEY", "").strip()
    if not key:
        raise RuntimeError(
            "OPENAI_API_KEY is not set. Please add it to your .env file "
            "or provide an API key in the node's Dynamic settings."
        )
    return OpenAIEmbeddings(api_key=key)


def _get_vectorstore(collection_name: str, api_key: str | None = None):
    from langchain_chroma import Chroma
    return Chroma(
        persist_directory=str(CHROMA_DIR),
        collection_name=collection_name,
        embedding_function=_get_embeddings(api_key),
    )


# ─── Source extractors ────────────────────────────────────────────────────────

def _load_pdf(source: str) -> list[Document]:
    from langchain_community.document_loaders import PyPDFLoader
    return PyPDFLoader(source).load()


def _load_docx(source: str) -> list[Document]:
    from langchain_community.document_loaders import Docx2txtLoader
    return Docx2txtLoader(source).load()


def _load_pptx(source: str) -> list[Document]:
    try:
        from pptx import Presentation  # type: ignore
    except ImportError:
        raise RuntimeError("python-pptx is required. Run: pip install python-pptx")
    prs = Presentation(source)
    texts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
    return [Document(page_content="\n\n".join(texts), metadata={"source": source})]


def _load_txt(source: str) -> list[Document]:
    from langchain_community.document_loaders import TextLoader
    return TextLoader(source, encoding="utf-8").load()


def _load_html(source: str) -> list[Document]:
    from langchain_community.document_loaders import UnstructuredHTMLLoader
    return UnstructuredHTMLLoader(source).load()


def _load_url(source: str) -> list[Document]:
    try:
        from newspaper import Article  # type: ignore
        article = Article(source)
        article.download()
        article.parse()
        text = article.text
    except Exception:
        try:
            from bs4 import BeautifulSoup  # type: ignore
            resp = requests.get(source, timeout=15)
            resp.raise_for_status()
            soup = BeautifulSoup(resp.text, "html.parser")
            for tag in soup(["script", "style", "nav", "footer"]):
                tag.decompose()
            text = soup.get_text(separator="\n", strip=True)
        except Exception as e:
            raise RuntimeError(f"Failed to fetch URL {source}: {e}")
    return [Document(page_content=text, metadata={"source": source})]


_EXT_LOADERS: dict[str, Any] = {
    ".pdf":  _load_pdf,
    ".docx": _load_docx,
    ".pptx": _load_pptx,
    ".txt":  _load_txt,
    ".md":   _load_txt,
    ".html": _load_html,
    ".htm":  _load_html,
}

_SOURCE_TYPE_LOADERS: dict[str, Any] = {
    "pdf":  _load_pdf,
    "docx": _load_docx,
    "pptx": _load_pptx,
    "txt":  _load_txt,
    "html": _load_html,
    "url":  _load_url,
}


# ─── Local file discovery ─────────────────────────────────────────────────────

def discover_local_files(path: str) -> list[str]:
    """
    Given a file path, list of paths (comma-separated), or directory path,
    return all supported file paths.
    """
    p = Path(path)
    found: list[str] = []

    if p.is_dir():
        for f in sorted(p.rglob("*")):
            if f.is_file() and f.suffix.lower() in SUPPORTED_EXTENSIONS:
                found.append(str(f))
    elif p.is_file():
        if p.suffix.lower() in SUPPORTED_EXTENSIONS:
            found.append(str(p))
    else:
        # Could be comma-separated list of paths
        parts = [s.strip() for s in path.split(",") if s.strip()]
        for part in parts:
            pp = Path(part)
            if pp.is_file() and pp.suffix.lower() in SUPPORTED_EXTENSIONS:
                found.append(str(pp))
            elif pp.is_dir():
                for f in sorted(pp.rglob("*")):
                    if f.is_file() and f.suffix.lower() in SUPPORTED_EXTENSIONS:
                        found.append(str(f))

    return found


def load_local_documents(path: str) -> list[Document]:
    """Load all supported documents from a local path/file/folder/list."""
    files = discover_local_files(path)
    if not files:
        raise ValueError(
            f"No supported files found at: {path!r}. "
            f"Supported types: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    docs: list[Document] = []
    for fp in files:
        ext = Path(fp).suffix.lower()
        loader_fn = _EXT_LOADERS.get(ext)
        if loader_fn:
            try:
                docs.extend(loader_fn(fp))
                print(f"  ✅ Loaded local file: {fp}")
            except Exception as e:
                print(f"  ⚠️  Failed to load {fp}: {e}")
    if not docs:
        raise ValueError(f"Could not extract content from any file in: {path!r}")
    return docs


# ─── Non-local source loading ─────────────────────────────────────────────────

def _detect_source_type(source: str, explicit_type: str | None) -> str:
    if explicit_type and explicit_type.lower() in _SOURCE_TYPE_LOADERS:
        return explicit_type.lower()
    if source.startswith("http://") or source.startswith("https://"):
        if source.lower().endswith(".pdf"):
            return "url_pdf"
        return "url"
    ext = Path(source).suffix.lower().lstrip(".")
    if ext in _SOURCE_TYPE_LOADERS:
        return ext
    raise ValueError(
        f"Cannot determine source type for: {source!r}. Set source_type explicitly."
    )


def load_nonlocal_documents(source: str, source_type: str | None = None) -> list[Document]:
    detected = _detect_source_type(source, source_type)

    if detected == "url_pdf":
        resp = requests.get(source, timeout=30)
        resp.raise_for_status()
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
            f.write(resp.content)
            tmp_path = f.name
        try:
            return _load_pdf(tmp_path)
        finally:
            os.unlink(tmp_path)

    loader_fn = _SOURCE_TYPE_LOADERS.get(detected)
    if not loader_fn:
        raise ValueError(f"Unsupported source type: {detected!r}")
    return loader_fn(source)


# ─── Chunking ─────────────────────────────────────────────────────────────────

def chunk_documents(
    docs: list[Document],
    chunk_size: int = 500,
    chunk_overlap: int = 100,
) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
    )
    return splitter.split_documents(docs)


# ─── Main retrieve entry point ────────────────────────────────────────────────

def retrieve(
    query: str,
    source: str,
    source_locality: str = "nonlocal",   # "local" | "nonlocal"
    source_type: str | None = None,
    chunk_size: int = 500,
    chunk_overlap: int = 100,
    k: int = 5,
    api_key: str | None = None,          # None -> .env (Static) | str -> Dynamic
) -> tuple[list[Document], list[float]]:
    """
    Full pipeline:
      1. Load documents (local or non-local)
      2. Chunk
      3. Embed → ChromaDB
      4. Retrieve top-k with confidence scores
    OpenAI is only touched here — never at import/startup time.
    """
    print(f"\n📂 Loading documents | locality={source_locality} | source={source!r}")

    if source_locality == "local":
        raw_docs = load_local_documents(source)
    else:
        raw_docs = load_nonlocal_documents(source, source_type)

    chunks = chunk_documents(raw_docs, chunk_size=chunk_size, chunk_overlap=chunk_overlap)
    print(f"📦 {len(chunks)} chunks from {len(raw_docs)} document(s)")

    safe_name = "".join(c if c.isalnum() else "_" for c in (source or "default"))[:60]
    vs = _get_vectorstore(safe_name, api_key)
    vs.add_documents(chunks)

    results = vs.similarity_search_with_score(query, k=k)
    docs_out = [doc for doc, _ in results]
    scores   = [round((1 / (1 + score)) * 100, 2) for _, score in results]
    return docs_out, scores